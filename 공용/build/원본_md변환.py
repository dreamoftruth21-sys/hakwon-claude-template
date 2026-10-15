"""<반>/학교자료 의 PDF·HWP·DOCX 텍스트를 <반>/학교자료/_텍스트/ 로 추출한다.

- 원본은 건드리지 않는다. 같은 하위 경로에 <파일명>.<확장자>.md 로 쓴다.
- 원본보다 md가 새것이면 건너뛴다(증분).
- 텍스트가 거의 안 나오는 PDF(스캔본)는 _스캔본목록.txt 에 적고 md는 만들지 않는다.

사용: py 공용/build/원본_md변환.py [반이름 ...]   (생략하면 전체 반)
"""
import sys, zlib, struct
from pathlib import Path

import pymupdf
import olefile

ROOT = Path(__file__).resolve().parents[2]
MIN_CHARS_PER_PAGE = 40


def pdf_text(path):
    doc = pymupdf.open(path)
    pages = [p.get_text("text").strip() for p in doc]
    n = len(pages)
    doc.close()
    total = sum(len(t) for t in pages)
    if n == 0 or total / n < MIN_CHARS_PER_PAGE:
        return None
    return "\n\n".join(f"<!-- p.{i} -->\n{t}" for i, t in enumerate(pages, 1) if t)


def hwp_text(path):
    ole = olefile.OleFileIO(str(path))
    header = ole.openstream("FileHeader").read()
    compressed = bool(header[36] & 1)
    sections = sorted(
        (e for e in ole.listdir() if e[0] == "BodyText"),
        key=lambda e: int(e[1].replace("Section", "")),
    )
    out = []
    for sec in sections:
        data = ole.openstream(sec).read()
        if compressed:
            data = zlib.decompress(data, -15)
        i = 0
        while i + 4 <= len(data):
            h = struct.unpack_from("<I", data, i)[0]
            tag, size = h & 0x3FF, (h >> 20) & 0xFFF
            i += 4
            if size == 0xFFF:
                size = struct.unpack_from("<I", data, i)[0]
                i += 4
            if tag == 67:  # HWPTAG_PARA_TEXT
                out.append(_hwp_para(data[i:i + size]))
            i += size
    ole.close()
    text = "\n".join(t for t in out if t.strip())
    return text or None


def _hwp_para(raw):
    chars, j = [], 0
    while j + 2 <= len(raw):
        c = struct.unpack_from("<H", raw, j)[0]
        if c in (1, 2, 3, 11, 12, 14, 15, 16, 17, 18, 21, 22, 23):
            j += 16  # 확장 컨트롤: 8 wchar
            continue
        if c in (4, 5, 6, 7, 8, 9, 19, 20):
            if c == 9:
                chars.append("\t")
            j += 16  # 인라인 컨트롤(탭 포함)
            continue
        if c in (10, 13):
            chars.append("\n")
        elif c >= 32:
            chars.append(chr(c))
        j += 2
    return "".join(chars).strip()


def docx_text(path):
    import docx
    d = docx.Document(str(path))
    lines = [p.text for p in d.paragraphs]
    for tb in d.tables:
        for row in tb.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    text = "\n".join(l for l in lines if l.strip())
    return text or None


HANDLERS = {".pdf": pdf_text, ".hwp": hwp_text, ".docx": docx_text}


def classes():
    """반 폴더 = 학원 폴더 바로 아래에서 CLAUDE.md 와 수업일지.md 를 함께 가진 폴더."""
    return sorted(p for p in ROOT.iterdir()
                  if p.is_dir() and (p / "CLAUDE.md").exists() and (p / "수업일지.md").exists())


def run(cls_dir):
    src_root = cls_dir / "학교자료"
    dst_root = cls_dir / "학교자료" / "_텍스트"
    if not src_root.exists():
        return
    stats = {"ok": 0, "skip": 0, "scan": 0, "fail": 0}
    scans, fails = [], []
    for src in sorted(src_root.rglob("*")):
        fn = HANDLERS.get(src.suffix.lower())
        if not fn or not src.is_file():
            continue
        rel = src.relative_to(src_root)
        dst = dst_root / rel.parent / (src.name + ".md")
        if dst.exists() and dst.stat().st_mtime >= src.stat().st_mtime:
            stats["skip"] += 1
            continue
        try:
            text = fn(src)
        except Exception as e:
            stats["fail"] += 1
            fails.append(f"{rel}\t{type(e).__name__}: {e}")
            continue
        if text is None:
            stats["scan"] += 1
            scans.append(str(rel))
            continue
        dst.parent.mkdir(parents=True, exist_ok=True)
        dst.write_text(f"# {src.name}\n원본: 학교자료/{rel.as_posix()}\n\n{text}\n", encoding="utf-8", errors="replace")
        stats["ok"] += 1
    dst_root.mkdir(exist_ok=True)
    (dst_root / "_스캔본목록.txt").write_text("\n".join(scans), encoding="utf-8")
    (dst_root / "_실패목록.txt").write_text("\n".join(fails), encoding="utf-8")
    print(cls_dir.name, stats)


if __name__ == "__main__":
    names = sys.argv[1:] or [p.name for p in classes()]
    for n in names:
        run(ROOT / n)
